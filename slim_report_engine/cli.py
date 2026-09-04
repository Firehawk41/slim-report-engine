"""Command-line entry point: reads a filled-in Testing Request workbook,
builds report content for every sample, writes the output workbook.

Designed to be invoked from a thin VBA trigger (per the project's
deployment decision — a familiar Excel button stays the UX for lab
technicians, but VBA no longer owns any content/formatting logic, just
the invocation). Reports success or failure via process exit code plus
a message -- printed to stdout/stderr for a normal console invocation,
AND written to --result-file when given, since the real VBA hook
launches this hidden (WScript.Shell.Run with a hidden window style, not
.Exec) specifically so no console window is ever shown to the
technician -- a hidden launch has no stdout/stderr pipes to read, so
--result-file is the only channel VBA actually has. See README.md's
"Invocation from VBA" section for the calling convention.

EXIT CODES:
  0  success -- output workbook written, its path is the message
  1  a recognized, actionable error (unresolved customer/chemical,
     unsupported analysis/DM5 chemical, empty submission, bad input
     file) -- the message is meant to be shown to the technician
     directly, e.g. via MsgBox
  2  an unexpected error (a real bug, not a data problem) -- the
     message is a full traceback, not meant for a technician to
     action
"""

from __future__ import annotations

import argparse
import sys
import traceback
from pathlib import Path

import openpyxl

from infrastructure.database import make_engine, make_session
from slim_domain.domain.analysis.analysis_service import AnalysisService
from slim_domain.domain.chemical.chemical_service import ChemicalService
from slim_domain.domain.customer.customer_service import CustomerService
from slim_domain.domain.element.element_service import ElementService
from slim_domain.domain.specification.specification_service import SpecificationService
from slim_domain.domain.tr.tr_form_input_resolver import (
    TRFormInputResolver,
    UnresolvedChemicalError,
    UnresolvedCustomerError,
)
from slim_domain.domain.tr.tr_submission_service import TRSubmissionService

from slim_report_engine.lab_identity import LAB_HEADER_LINES
from slim_report_engine.reporting import report_writer
from slim_report_engine.reporting.submission_report_builder import build_submission_sheets


class ReportEngineError(Exception):
    """A recognized, actionable failure. Message is safe to show a technician directly."""


def run(input_path: Path, output_path: Path, db_url: str | None = None) -> Path:
    engine = make_engine(db_url)
    session = make_session(engine)

    chemical_svc = ChemicalService(session)
    analysis_svc = AnalysisService(session)
    element_svc = ElementService(session)
    customer_svc = CustomerService(session)
    specification_svc = SpecificationService(session)
    resolver = TRFormInputResolver(customer_svc, chemical_svc)
    submission_svc = TRSubmissionService(session, chemical_svc, analysis_svc, element_svc, resolver)

    try:
        submission = submission_svc.build_from_file(input_path)
    except (UnresolvedCustomerError, UnresolvedChemicalError) as e:
        raise ReportEngineError(
            f"{e} -- this needs to be added to the database before a report can be generated for it."
        ) from e
    except ValueError as e:
        raise ReportEngineError(f"could not read {input_path.name}: {e}") from e

    customer = customer_svc.load_customer(submission.customer_id)
    if customer is None:
        raise ReportEngineError(
            f"customer_id {submission.customer_id} was resolved during parsing but no longer loads "
            "-- this indicates a real bug, not a data problem."
        )

    try:
        sheets = build_submission_sheets(
            submission, customer, chemical_svc, analysis_svc, element_svc, specification_svc
        )
    except ValueError as e:
        raise ReportEngineError(str(e)) from e

    if not sheets:
        raise ReportEngineError(f"no samples found in {input_path.name} -- nothing to write.")

    # Confirmed real (every completed report checked): the filled-in
    # Testing Request form itself is appended as the LAST tab of the
    # output workbook. Building on top of the loaded input workbook
    # (rather than a fresh one) keeps that sheet's real formatting,
    # formulas, and dropdowns exactly as submitted -- no manual
    # cell-by-cell copy needed, which openpyxl has no built-in support
    # for across two separate workbook objects anyway.
    wb = openpyxl.load_workbook(input_path)
    form_sheet_name = _find_form_sheet_name(wb)
    for sheet in sheets:
        ws = wb.create_sheet(sheet.name)
        report_writer.apply_zoom(ws, sheet.zoom)
        report_writer.apply_column_widths(ws, sheet.column_widths)
        report_writer.write_sections(ws, list(sheet.sections), start_row=1)
        report_writer.apply_header_footer(
            ws,
            center_header=sheet.header_title,
            customer_name=customer.name,
            customer_address_line1=customer.street_address,
            customer_address_line2=customer.city,
            lab_header_lines=LAB_HEADER_LINES,
        )
        for cell_address, value in sheet.extra_cell_stamps:
            ws[cell_address] = value
        if sheet.processing_time_stamp is not None:
            report_writer.apply_processing_time_stamp(ws, *sheet.processing_time_stamp)

    if form_sheet_name is not None:
        wb.move_sheet(form_sheet_name, offset=len(wb.sheetnames))

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path


def _find_form_sheet_name(wb: openpyxl.Workbook) -> str | None:
    """Same detection rule as slim-domain's TRSubmissionService (the first
    sheet whose D3 cell contains "Testing Request Form") -- reimplemented
    here rather than imported since it needs a real (non-read-only)
    Worksheet object to check the same file build_from_file already
    parsed in read-only mode."""
    for ws in wb.worksheets:
        if "Testing Request Form" in str(ws["D3"].value or ""):
            return ws.title
    return None


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="Generate a blank lab report workbook from a Testing Request workbook."
    )
    parser.add_argument("input", type=Path, help="path to the filled-in Testing Request .xlsx")
    parser.add_argument(
        "-o", "--output", type=Path, default=None,
        help="output .xlsx path (default: <input>_report.xlsx, next to the input file)",
    )
    parser.add_argument(
        "--db", type=str, default=None,
        help="DB_URL override (default: .env's DB_URL, or sqlite:///report_engine.db)",
    )
    parser.add_argument(
        "--result-file", type=Path, default=None,
        help="write '<exit code>\\n<message>' here -- the only channel a hidden "
             "(no console window) launch has, since it has no stdout/stderr to read",
    )
    args = parser.parse_args(argv)

    if not args.input.exists():
        return _finish(1, f"input file not found: {args.input}", args.result_file)

    output_path = args.output or args.input.with_name(args.input.stem + "_report.xlsx")

    try:
        result_path = run(args.input, output_path, args.db)
    except ReportEngineError as e:
        return _finish(1, str(e), args.result_file)
    except Exception:
        return _finish(2, traceback.format_exc(), args.result_file)

    return _finish(0, str(result_path), args.result_file)


def _finish(code: int, message: str, result_file: Path | None) -> int:
    """Reports (code, message) on stdout/stderr (guarded: a --windowed-style
    frozen build with no console attached has sys.stdout/sys.stderr as None,
    though the real deployed build keeps a console -- just a hidden one) and,
    if given, writes it to result_file for the hidden-launch VBA hook to read.
    Written via a same-directory temp file + os.replace (atomic on Windows)
    so a concurrent reader never observes a partially-written file.
    """
    stream = sys.stdout if code == 0 else sys.stderr
    if stream is not None:
        print(message, file=stream)
    if result_file is not None:
        tmp = result_file.with_name(result_file.name + ".tmp")
        tmp.write_text(f"{code}\n{message}", encoding="utf-8")
        tmp.replace(result_file)
    return code


if __name__ == "__main__":
    raise SystemExit(main())
