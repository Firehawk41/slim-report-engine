"""End-to-end test of the actual CLI entry point (slim_report_engine.cli)
against a real (temp-file) SQLite database and the real blank Chemical
intake form -- the same contract a VBA trigger would rely on: exit code
0 + output file on success, exit code 1 + a clear stderr message on a
recognized data problem (e.g. an unresolved customer).
"""

import os
from datetime import date
from pathlib import Path

import openpyxl
import pytest
from sqlalchemy import create_engine, text
from sqlalchemy.orm import Session

from infrastructure.database import Base
from slim_domain.domain.analysis.analysis_repository import _AnalysisRow, _FormAnalysisRow
from slim_domain.domain.chemical.chemical_repository import _ChemicalRow, _FormChemicalRow
from slim_domain.domain.customer.customer_repository import _CustomerRow, _FormCustomerRow
from slim_domain.domain.element.element_repository import _ElementRow

from slim_report_engine.cli import main
from slim_report_engine.lab_identity import LAB_HEADER_LINES
from slim_report_engine.reporting.presets import analyte_presets

_FORM_PATH = (
    Path(os.environ.get("PVB_BRIDGE_PATH", r"C:\Users\Jamie\Python-VBA-Bridge"))
    / "ReportCreator Project"
    / "Doc-F31 PRECILAB Analytical Testing Request Form - Chemicals.xlsx"
)


def _seed(db_url: str) -> None:
    engine = create_engine(db_url)
    Base.metadata.create_all(engine)
    with Session(engine) as s:
        for symbol in analyte_presets.trace_elements_36():
            s.add(_ElementRow(element_symbol=symbol, element_name=symbol))

        customer = _CustomerRow(
            customer_name="Acme Corp", street_address="123 Main St",
            city="Springfield", state="IL", postal_code="62701", country="USA",
        )
        s.add(customer)
        s.flush()
        s.add(_FormCustomerRow(
            form_customer="Acme Corp", form_address="123 Main St", form_address_2="Springfield, IL",
            customer_id=customer.id,
        ))

        chemical = _ChemicalRow(chemical_name="Test Acid", metals_prep="N/A", silicon_prep="N/A", ions_prep="N/A")
        s.add(chemical)
        s.flush()
        s.add(_FormChemicalRow(form_name="Test Acid Matrix", chemical_id=chemical.id))

        analysis = _AnalysisRow(analysis_name="36 Elements", analysis_description="36 elements (by ICP-MS)")
        s.add(analysis)
        s.flush()
        s.add(_FormAnalysisRow(form_name="36 Elements", analysis_id=analysis.id))

        s.execute(text(
            "CREATE TABLE invoice_default_contacts "
            "(contact_email TEXT, customer_id INTEGER, is_default INTEGER, is_active INTEGER)"
        ))
        s.commit()


def _make_input_workbook(path: Path, customer_name: str = "Acme Corp") -> None:
    wb = openpyxl.load_workbook(_FORM_PATH)
    ws = wb.active
    ws["C10"] = customer_name
    ws["C11"] = "123 Main St"
    ws["C12"] = "Springfield, IL"
    ws["P5"] = date(2026, 3, 5)
    ws["B22"] = "S-001"
    ws["C22"] = "Test Acid Matrix"
    ws["D22"] = "36 Elements"
    ws["K22"] = "Next Day"
    wb.save(path)


@pytest.mark.skipif(not _FORM_PATH.exists(), reason="real Chemical intake form not available in this environment")
def test_cli_success_writes_output_and_exits_zero(tmp_path, capsys):
    db_path = tmp_path / "test.db"
    db_url = f"sqlite:///{db_path}"
    _seed(db_url)

    input_path = tmp_path / "input.xlsx"
    _make_input_workbook(input_path)
    output_path = tmp_path / "output.xlsx"

    exit_code = main(["--db", db_url, str(input_path), "-o", str(output_path)])

    assert exit_code == 0
    assert output_path.exists()

    out, err = capsys.readouterr()
    assert str(output_path) in out
    assert err == ""

    wb = openpyxl.load_workbook(output_path)
    assert "S-001" in wb.sheetnames
    ws = wb["S-001"]
    assert ws.oddHeader.center.text == "Test Acid Matrix"
    assert ws.oddHeader.left.text == "\n".join(LAB_HEADER_LINES)
    assert "Acme Corp" in ws.oddHeader.right.text
    assert ws["A1"].value == "Test Acid Matrix"
    # The standard "mmddyy-Chemical-Customer-SampleID" sample-ID string,
    # stamped into the section's own "Sample Identification:" row.
    assert ws["D1"].value == "030526-Test Acid Matrix-Acme Corp-S-001"


@pytest.mark.skipif(not _FORM_PATH.exists(), reason="real Chemical intake form not available in this environment")
def test_cli_unresolved_customer_exits_one_with_clear_message(tmp_path, capsys):
    db_path = tmp_path / "test.db"
    db_url = f"sqlite:///{db_path}"
    _seed(db_url)

    input_path = tmp_path / "input.xlsx"
    _make_input_workbook(input_path, customer_name="Totally Unknown Customer")
    output_path = tmp_path / "output.xlsx"

    exit_code = main(["--db", db_url, str(input_path), "-o", str(output_path)])

    assert exit_code == 1
    assert not output_path.exists()

    out, err = capsys.readouterr()
    assert "Totally Unknown Customer" in err
    assert "database" in err.lower()


def test_cli_missing_input_file_exits_one(tmp_path, capsys):
    exit_code = main([str(tmp_path / "does_not_exist.xlsx")])
    assert exit_code == 1
    out, err = capsys.readouterr()
    assert "not found" in err
