"""End-to-end proof that the real pieces actually fit together: read a
real (blank, then filled-in-memory) Chemical intake form with
slim-domain's real TRSubmissionFormParser and real Analysis/Element
services backed by a real (in-memory) SQLite database, then feed the
resulting TRSample straight into this project's own Transform-stage
orchestrator with no adapter code in between.

Not a substitute for the unit tests elsewhere in this package (which
cover many more shapes/edge cases with fakes) -- this proves the actual
wiring, not just each piece in isolation.
"""

import os
from pathlib import Path

import openpyxl
import pytest
from sqlalchemy import create_engine, text
from sqlalchemy.orm import Session

from infrastructure.database import Base
from slim_domain.domain.analysis.analysis_repository import _AnalysisRow, _FormAnalysisRow
from slim_domain.domain.analysis.analysis_service import AnalysisService
from slim_domain.domain.chemical.chemical import Chemical
from slim_domain.domain.chemical.chemical_service import ChemicalService
from slim_domain.domain.customer.customer import Customer
from slim_domain.domain.element.element_repository import _ElementRow
from slim_domain.domain.element.element_service import ElementService
from slim_domain.domain.tr.tr_submission_form_parser import TRSubmissionFormParser

from slim_report_engine.reporting import chemical_water_report_builder
from slim_report_engine.reporting.presets import analyte_presets

_FORM_PATH = (
    Path(os.environ.get("PVB_BRIDGE_PATH", r"C:\Users\Jamie\Python-VBA-Bridge"))
    / "ReportCreator Project"
    / "Doc-F31 PRECILAB Analytical Testing Request Form - Chemicals.xlsx"
)

_STUB_CUSTOMER = Customer(
    id=1, name="Acme Corp", street_address="123 Main St",
    city="Springfield", state="IL", postal_code="62701", country="USA",
)
_STUB_CHEMICAL = Chemical(id=7, name="Test Acid", metals_prep="N/A", silicon_prep="N/A", ions_prep="N/A")


class _StubResolver:
    def resolve_customer(self, name, address, address2):
        return _STUB_CUSTOMER

    def resolve_chemical(self, form_name):
        return _STUB_CHEMICAL


@pytest.fixture
def session():
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    with Session(engine) as s:
        for symbol in analyte_presets.trace_elements_36():
            s.add(_ElementRow(element_symbol=symbol, element_name=symbol))
        s.add(_AnalysisRow(analysis_name="36 Elements", analysis_description="36 elements (by ICP-MS)"))
        s.commit()
        analysis_id = s.query(_AnalysisRow).filter_by(analysis_name="36 Elements").one().id
        s.add(_FormAnalysisRow(form_name="36 Elements", analysis_id=analysis_id))
        s.execute(text(
            "CREATE TABLE invoice_default_contacts "
            "(contact_email TEXT, customer_id INTEGER, is_default INTEGER, is_active INTEGER)"
        ))
        s.commit()
        yield s


@pytest.mark.skipif(not _FORM_PATH.exists(), reason="real Chemical intake form not available in this environment")
def test_real_chemical_form_flows_through_parse_and_transform(session):
    wb = openpyxl.load_workbook(_FORM_PATH)
    ws = wb.active

    # Fill in the fixed customer cells (StubResolver ignores their content,
    # but a real resolver would need them -- filled for realism).
    ws["C10"] = "Acme Corp"
    ws["C11"] = "123 Main St"
    ws["C12"] = "Springfield, IL"

    # One sample row. Real blank form's B20 == "Sample ID" -> first data
    # row is 22 (confirmed against the real form earlier this session).
    ws["B22"] = "S-001"
    ws["C22"] = "Test Acid Matrix"
    ws["D22"] = "36 Elements"
    ws["K22"] = "Next Day"  # processing time, column 11 for Chemical

    analysis_svc = AnalysisService(session)
    chemical_svc = ChemicalService(session)
    element_svc = ElementService(session)
    resolver = _StubResolver()

    parser = TRSubmissionFormParser(
        session=session,
        chemical_service=chemical_svc,
        analysis_service=analysis_svc,
        element_service=element_svc,
        input_resolver=resolver,
    )

    submission = parser.build_from_worksheet(ws, filename="test.xlsx")

    assert len(submission.samples) == 1
    sample = submission.samples[0]
    assert sample.sample_name == "S-001"

    resolved_names = [analysis_svc.load_analysis(aid).name for aid in sample.analysis_ids]
    assert "36 Elements" in resolved_names

    sections = chemical_water_report_builder.build_sections(sample, analysis_svc, element_svc)

    assert len(sections) == 1
    section = sections[0]
    assert section.id == "36 Elements"
    data_rows = [r for r in section.rows if r.style_name == "DataLabel" and r.get_value(3)]
    assert len(data_rows) == 36
    avg_row = next(r for r in section.rows if r.get_value(2) and "AVERAGE" in str(r.get_value(2)))
    assert avg_row.get_value(2) == "AVERAGE / 36 Tr.Elts"
