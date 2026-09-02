"""End-to-end proof for the Wafer path, mirroring
test_chemical_water_pipeline.py: read the real blank Wafer intake form
with slim-domain's real TRSubmissionFormParser (backed by a real
in-memory SQLite session and real Analysis/Element services), fill in two
sample rows sharing one grouping combination, and feed the resulting
TRSubmission straight into wafer_submission_builder.build_wafer_sheets
with no adapter code in between.
"""

import os
from datetime import date
from pathlib import Path

import openpyxl
import pytest
from sqlalchemy import create_engine, text
from sqlalchemy.orm import Session

from infrastructure.database import Base
from slim_domain.domain.analysis.analysis_repository import _AnalysisRow, _FormAnalysisRow
from slim_domain.domain.analysis.analysis_service import AnalysisService
from slim_domain.domain.chemical.chemical import Chemical
from slim_domain.domain.chemical.chemical_service import ChemicalService
from slim_domain.domain.customer.customer import Customer
from slim_domain.domain.element.element_repository import _ElementRow
from slim_domain.domain.element.element_service import ElementService
from slim_domain.domain.tr.tr_submission_form_parser import TRSubmissionFormParser

from slim_report_engine.reporting import wafer_submission_builder
from slim_report_engine.reporting.presets import analyte_presets

_FORM_PATH = (
    Path(os.environ.get("PVB_BRIDGE_PATH", r"C:\Users\Jamie\Python-VBA-Bridge"))
    / "ReportCreator Project"
    / "Doc-F35 PRECILAB Analytical Testing Request Form - Wafers.xlsx"
)

_STUB_CUSTOMER = Customer(
    id=1, name="Acme Corp", street_address="123 Main St",
    city="Springfield", state="IL", postal_code="62701", country="USA",
)
_STUB_CHEMICAL = Chemical(id=7, name="N/A", metals_prep="N/A", silicon_prep="N/A", ions_prep="N/A")


class _StubResolver:
    def resolve_customer(self, name, address, address2):
        return _STUB_CUSTOMER

    def resolve_chemical(self, form_name):
        return _STUB_CHEMICAL


@pytest.fixture
def session():
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    with Session(engine) as s:
        for symbol in analyte_presets.trace_elements_36():
            s.add(_ElementRow(element_symbol=symbol, element_name=symbol))
        s.add(_AnalysisRow(analysis_name="36 Elements", analysis_description="36 elements (by ICP-MS)"))
        s.commit()
        analysis_id = s.query(_AnalysisRow).filter_by(analysis_name="36 Elements").one().id
        s.add(_FormAnalysisRow(form_name="36 Elements", analysis_id=analysis_id))
        s.execute(text(
            "CREATE TABLE invoice_default_contacts "
            "(contact_email TEXT, customer_id INTEGER, is_default INTEGER, is_active INTEGER)"
        ))
        s.commit()
        yield s


@pytest.mark.skipif(not _FORM_PATH.exists(), reason="real Wafer intake form not available in this environment")
def test_real_wafer_form_flows_through_parse_and_group_and_transform(session):
    wb = openpyxl.load_workbook(_FORM_PATH)
    ws = wb.active

    ws["C10"] = "Acme Corp"
    ws["C11"] = "123 Main St"
    ws["C12"] = "Springfield, IL"
    ws["L5"] = date(2026, 3, 5)  # Date received (Wafer column map)

    # Two sample rows sharing the same (Reporting Units, Wafer Size,
    # Element selection, Processing Time) -- should land on one sheet.
    # Real form's B20 == "Sample ID / Slot Number" -> first data row is 22
    # (confirmed via the _get_sample_range fix this session).
    for row, slot_label in ((22, "Slot-1"), (23, "Slot-2")):
        ws.cell(row=row, column=2).value = slot_label   # Sample ID / Slot Number
        ws.cell(row=row, column=3).value = "atoms/cm^2"  # Reporting Unit
        ws.cell(row=row, column=4).value = "150mm"       # Wafer Size
        ws.cell(row=row, column=5).value = "36 Elements"  # Number of Elements
        ws.cell(row=row, column=9).value = "Next Day"    # Processing Time

    analysis_svc = AnalysisService(session)
    chemical_svc = ChemicalService(session)
    element_svc = ElementService(session)
    resolver = _StubResolver()

    parser = TRSubmissionFormParser(
        session=session,
        chemical_service=chemical_svc,
        analysis_service=analysis_svc,
        element_service=element_svc,
        input_resolver=resolver,
    )

    submission = parser.build_from_worksheet(ws, filename="test.xlsx")
    assert len(submission.samples) == 2

    results = wafer_submission_builder.build_wafer_sheets(
        submission, _STUB_CUSTOMER, analysis_svc, element_svc
    )

    assert len(results) == 1
    result = results[0]
    assert result.sheet_title == "150mm Wafers"
    assert result.apply_atoms_superscript_quirk is True

    sample_id_row = result.section.rows[5]
    assert sample_id_row.get_value(4) == f"{submission.date_received.strftime('%m%d%y')}-150mm Wafers-Acme Corp-Process Blank"
    assert sample_id_row.get_value(5) == f"{submission.date_received.strftime('%m%d%y')}-150mm Wafers-Acme Corp-Slot-1"
    assert sample_id_row.get_value(6) == f"{submission.date_received.strftime('%m%d%y')}-150mm Wafers-Acme Corp-Slot-2"

    data_rows = [r for r in result.section.rows if r.style_name == "DataLabel"]
    assert len(data_rows) == 36
