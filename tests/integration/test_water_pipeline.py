"""End-to-end proof for the Water path, mirroring
test_chemical_water_pipeline.py: Water has its own real intake form with a
genuinely different column layout from Chemical (see DOMAIN_BRIEF.md's
column map) -- this confirms the real parser handles that layout
correctly too, not just Chemical's.
"""

import os
from pathlib import Path

import openpyxl
import pytest
from sqlalchemy import create_engine, text
from sqlalchemy.orm import Session

from infrastructure.database import Base
from slim_domain.domain.analysis.analysis_repository import _AnalysisRow, _FormAnalysisRow
from slim_domain.domain.analysis.analysis_service import AnalysisService
from slim_domain.domain.chemical.chemical import Chemical
from slim_domain.domain.chemical.chemical_service import ChemicalService
from slim_domain.domain.customer.customer import Customer
from slim_domain.domain.element.element_service import ElementService
from slim_domain.domain.tr.tr_submission_form_parser import TRSubmissionFormParser

from slim_report_engine.reporting import chemical_water_report_builder

_FORM_PATH = (
    Path(os.environ.get("PVB_BRIDGE_PATH", r"C:\Users\Jamie\Python-VBA-Bridge"))
    / "ReportCreator Project"
    / "Doc-F34 PRECILAB Analytical Testing Request Form - Water.xlsx"
)

_STUB_CUSTOMER = Customer(
    id=1, name="Acme Corp", street_address="123 Main St",
    city="Springfield", state="IL", postal_code="62701", country="USA",
)
_STUB_CHEMICAL = Chemical(id=2, name="Water", metals_prep="N/A", silicon_prep="N/A", ions_prep="N/A")


class _StubResolver:
    def resolve_customer(self, name, address, address2):
        return _STUB_CUSTOMER

    def resolve_chemical(self, form_name):
        return _STUB_CHEMICAL


@pytest.fixture
def session():
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    with Session(engine) as s:
        s.add(_AnalysisRow(analysis_name="TOC", analysis_description="Total Organic Carbon"))
        s.commit()
        analysis_id = s.query(_AnalysisRow).filter_by(analysis_name="TOC").one().id
        s.add(_FormAnalysisRow(form_name="TOC", analysis_id=analysis_id))
        s.execute(text(
            "CREATE TABLE invoice_default_contacts "
            "(contact_email TEXT, customer_id INTEGER, is_default INTEGER, is_active INTEGER)"
        ))
        s.commit()
        yield s


@pytest.mark.skipif(not _FORM_PATH.exists(), reason="real Water intake form not available in this environment")
def test_real_water_form_flows_through_parse_and_transform(session):
    wb = openpyxl.load_workbook(_FORM_PATH)
    ws = wb.active

    ws["C10"] = "Acme Corp"
    ws["C11"] = "123 Main St"
    ws["C12"] = "Springfield, IL"

    # Water's real column map (confirmed against the real form): sample
    # table starts col 2, "chemical" is the fixed literal "Water" (never
    # read from a cell, per DOMAIN_BRIEF.md), analysis columns start col
    # 4, processing time col 14. B20 == "Sample ID" exactly (unlike
    # Wafer) -> first_row = 22.
    ws["B22"] = "S-001"
    ws["D22"] = "TOC"
    ws["N22"] = "Next Day"  # processing time, column 14 for Water

    analysis_svc = AnalysisService(session)
    chemical_svc = ChemicalService(session)
    element_svc = ElementService(session)
    resolver = _StubResolver()

    parser = TRSubmissionFormParser(
        session=session,
        chemical_service=chemical_svc,
        analysis_service=analysis_svc,
        element_service=element_svc,
        input_resolver=resolver,
    )

    submission = parser.build_from_worksheet(ws, filename="test.xlsx")

    assert len(submission.samples) == 1
    sample = submission.samples[0]
    assert sample.form_chemical_name == "Water"  # constant, never read from a cell

    resolved_names = [analysis_svc.load_analysis(aid).name for aid in sample.analysis_ids]
    assert resolved_names == ["TOC"]

    sections = chemical_water_report_builder.build_sections(sample, analysis_svc, element_svc)

    assert len(sections) == 1
    section = sections[0]
    assert section.id == "TOC"
    data_row = next(r for r in section.rows if r.style_name == "DataLabel")
    assert data_row.get_value(2) == "TOC"
    assert data_row.get_value(3) == "ppb"
