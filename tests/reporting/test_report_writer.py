import openpyxl
import pytest

from slim_report_engine.reporting import report_writer
from slim_report_engine.reporting.report_row import ReportRow
from slim_report_engine.reporting.report_section import ReportSection


@pytest.fixture
def ws():
    wb = openpyxl.Workbook()
    return wb.active


# ---------------------------------------------------------------------------
# apply_style
# ---------------------------------------------------------------------------

def test_normal_style_leaves_default_formatting(ws):
    report_writer.apply_style(ws, row_index=1, min_col=1, max_col=3, style_name="Normal")
    cell = ws.cell(row=1, column=1)
    assert cell.font.bold is not True
    assert cell.border.left.style is None


def test_section_title_style_is_bold_no_border_no_fill(ws):
    """DM5's chemical-label title row (e.g. "NH4OH") -- confirmed real:
    bold, no border, no fill."""
    report_writer.apply_style(ws, row_index=1, min_col=1, max_col=5, style_name="SectionTitle")
    cell = ws.cell(row=1, column=1)
    assert cell.font.bold is True
    assert cell.border.left.style is None
    assert cell.fill.patternType is None


def test_sample_id_echo_wide_style(ws):
    """Confirmed real (Report Creator Template.xlsx, DM5-N/S, Wafers):
    column 1 = white fill + bold + border; columns 2/3 = pale cyan;
    column 4+ = pale cyan + bold, extending across every Wafer sample
    column."""
    report_writer.apply_style(ws, row_index=1, min_col=1, max_col=6, style_name="SampleIdEchoWide")
    col1 = ws.cell(row=1, column=1)
    assert col1.font.bold is True
    assert col1.fill.fgColor.rgb == "FFFFFFFF"
    assert col1.border.left.style == "thin"

    col2 = ws.cell(row=1, column=2)
    assert col2.fill.fgColor.rgb == "FFCCFFFF"
    assert col2.border.left.style == "thin"

    col4 = ws.cell(row=1, column=4)
    assert col4.fill.fgColor.rgb == "FFCCFFFF"
    assert col4.font.bold is True
    assert col4.alignment.horizontal == "center"
    assert col4.alignment.wrap_text is True  # confirmed real (D6 in the real template)


def test_sample_id_echo_simple_style_has_no_column_1_border():
    """Confirmed real: TOC/Electrical/Misc Analysis's Sample-ID row has NO
    border or fill at all in column 1, unlike the Wide variant."""
    wb = openpyxl.Workbook()
    ws = wb.active
    report_writer.apply_style(ws, row_index=1, min_col=1, max_col=4, style_name="SampleIdEchoSimple")
    col1 = ws.cell(row=1, column=1)
    assert col1.border.left.style is None
    assert col1.fill.patternType is None
    col4 = ws.cell(row=1, column=4)
    assert col4.fill.fgColor.rgb == "FFCCFFFF"


def test_sample_id_echo_assay_style_only_styles_columns_2_and_3():
    """DM5's Assay-shape "Sample #" row -- confirmed real: shifted one
    column left of the Wide variant, nothing styled outside columns 2/3."""
    wb = openpyxl.Workbook()
    ws = wb.active
    report_writer.apply_style(ws, row_index=1, min_col=1, max_col=4, style_name="SampleIdEchoAssay")
    assert ws.cell(row=1, column=1).border.left.style is None
    assert ws.cell(row=1, column=2).fill.fgColor.rgb == "FFCCFFFF"
    assert ws.cell(row=1, column=3).fill.fgColor.rgb == "FFCCFFFF"
    assert ws.cell(row=1, column=4).border.left.style is None


def test_column_header_style_not_bold_white_then_light_blue(ws):
    """Confirmed real: category label (col 2/3) = white, NOT bold;
    Results/Recovery/MDL (col 4+) = light blue, also NOT bold -- despite
    the style being named after modReportStyles.bas's "HeaderBold"."""
    report_writer.apply_style(ws, row_index=2, min_col=1, max_col=6, style_name="ColumnHeader")
    col2 = ws.cell(row=2, column=2)
    assert col2.font.bold is not True
    assert col2.fill.fgColor.rgb == "FFFFFFFF"
    col4 = ws.cell(row=2, column=4)
    assert col4.font.bold is not True
    assert col4.fill.fgColor.rgb == "FF99CCFF"
    assert col4.alignment.horizontal == "center"


def test_column_header_style_column_1_gets_avantgarde_font_no_fill():
    """DM5's element-panel header row has REAL "Specification" text in
    column 1 (AvantGarde, no fill) -- generic panels leave it blank, but
    applying the same font there is harmless since it's invisible."""
    wb = openpyxl.Workbook()
    ws = wb.active
    report_writer.apply_style(ws, row_index=2, min_col=1, max_col=6, style_name="ColumnHeader")
    col1 = ws.cell(row=2, column=1)
    assert col1.font.name == "AvantGarde"
    assert col1.fill.patternType is None
    assert col1.border.left.style == "thin"


def test_column_header_assay_style_results_start_at_column_3():
    """DM5's Assay-shape header row -- confirmed real: results (bold,
    light blue) start one column earlier than the generic ColumnHeader
    style, and columns 1/2 have no fill at all."""
    wb = openpyxl.Workbook()
    ws = wb.active
    report_writer.apply_style(ws, row_index=5, min_col=1, max_col=4, style_name="ColumnHeaderAssay")
    assert ws.cell(row=5, column=1).fill.patternType is None
    assert ws.cell(row=5, column=2).fill.patternType is None
    col3 = ws.cell(row=5, column=3)
    assert col3.fill.fgColor.rgb == "FF99CCFF"
    assert col3.font.bold is True


def test_data_label_style(ws):
    report_writer.apply_style(ws, row_index=2, min_col=1, max_col=1, style_name="DataLabel")
    cell = ws.cell(row=2, column=1)
    assert cell.border.left.style == "thin"
    assert cell.alignment.horizontal == "center"  # column 1 (spec-value convention)


def test_data_label_style_column_2_plus_is_left_aligned():
    wb = openpyxl.Workbook()
    ws = wb.active
    report_writer.apply_style(ws, row_index=2, min_col=1, max_col=3, style_name="DataLabel")
    cell = ws.cell(row=2, column=2)
    assert cell.alignment.horizontal == "left"
    assert cell.border.left.style == "thin"


def test_data_value_style(ws):
    report_writer.apply_style(ws, row_index=2, min_col=1, max_col=1, style_name="DataValue")
    cell = ws.cell(row=2, column=1)
    assert cell.border.left.style == "thin"
    assert cell.alignment.horizontal == "right"


def test_summary_row_style_pale_cyan_fill_skips_column_1():
    """The metals panel's AVERAGE/TOTAL rows -- confirmed real: same pale
    cyan fill as the Sample-ID row; no real template populates column 1
    here, so it's left completely unstyled."""
    wb = openpyxl.Workbook()
    ws = wb.active
    report_writer.apply_style(ws, row_index=40, min_col=1, max_col=6, style_name="SummaryRow")
    assert ws.cell(row=40, column=1).border.left.style is None
    col2 = ws.cell(row=40, column=2)
    assert col2.fill.fgColor.rgb == "FFCCFFFF"
    assert col2.border.left.style == "thin"


def test_normal_bold_style_no_border_no_fill_but_bold():
    """Wafer's "Notes:" and units-note rows."""
    wb = openpyxl.Workbook()
    ws = wb.active
    report_writer.apply_style(ws, row_index=3, min_col=1, max_col=1, style_name="NormalBold")
    cell = ws.cell(row=3, column=1)
    assert cell.font.bold is True
    assert cell.border.left.style is None
    assert cell.fill.patternType is None


def test_normal_bold_blue_style_bold_and_blue():
    """Confirmed real (Report Creator Template.xlsx row 5): the generic
    preamble's own "Blue font indicates..." line."""
    wb = openpyxl.Workbook()
    ws = wb.active
    report_writer.apply_style(ws, row_index=3, min_col=1, max_col=1, style_name="NormalBoldBlue")
    cell = ws.cell(row=3, column=1)
    assert cell.font.bold is True
    assert cell.font.color.rgb == "FF0000FF"
    assert cell.border.left.style is None
    assert cell.fill.patternType is None


def test_normal_blue_style_blue_not_bold():
    """Confirmed real (Wafers Report Template.xlsx row 4): Wafer's own
    equivalent line is blue but NOT bold, unlike the generic preamble's."""
    wb = openpyxl.Workbook()
    ws = wb.active
    report_writer.apply_style(ws, row_index=3, min_col=1, max_col=1, style_name="NormalBlue")
    cell = ws.cell(row=3, column=1)
    assert cell.font.bold is not True
    assert cell.font.color.rgb == "FF0000FF"


def test_unknown_style_raises(ws):
    with pytest.raises(ValueError, match="unknown style name"):
        report_writer.apply_style(ws, row_index=1, min_col=1, max_col=1, style_name="Nonexistent")


# ---------------------------------------------------------------------------
# apply_zoom
# ---------------------------------------------------------------------------

def test_apply_zoom_sets_zoom_scale(ws):
    report_writer.apply_zoom(ws, 90)
    assert ws.sheet_view.zoomScale == 90


def test_apply_zoom_wafer_value(ws):
    report_writer.apply_zoom(ws, 95)
    assert ws.sheet_view.zoomScale == 95


# ---------------------------------------------------------------------------
# write_row
# ---------------------------------------------------------------------------

def test_write_row_writes_values(ws):
    row = ReportRow(style_name="Normal")
    row.set_value(2, "Aluminum")
    row.set_value(3, "Al")
    report_writer.write_row(ws, row, row_index=5)
    assert ws.cell(row=5, column=2).value == "Aluminum"
    assert ws.cell(row=5, column=3).value == "Al"


def test_write_row_translates_formula_to_real_a1_range(ws):
    row = ReportRow(style_name="DataLabel")
    row.set_formula(4, "=AVERAGE(R[-4]C:R[-2]C)")
    report_writer.write_row(ws, row, row_index=40)
    assert ws.cell(row=40, column=4).value == "=AVERAGE(D36:D38)"


def test_write_row_applies_style(ws):
    row = ReportRow(style_name="ColumnHeader")
    row.set_value(2, "Element")
    report_writer.write_row(ws, row, row_index=1)
    cell = ws.cell(row=1, column=2)
    assert cell.font.bold is not True
    assert cell.fill.fgColor.rgb == "FFFFFFFF"


def test_write_row_applies_row_height_when_set(ws):
    row = ReportRow(style_name="Normal", row_height=60.0)
    report_writer.write_row(ws, row, row_index=3)
    assert ws.row_dimensions[3].height == 60.0


def test_write_row_leaves_row_height_unset_by_default(ws):
    row = ReportRow(style_name="Normal")
    report_writer.write_row(ws, row, row_index=3)
    assert ws.row_dimensions[3].height is None


def test_write_row_applies_merges(ws):
    row = ReportRow(style_name="Normal")
    row.add_merge(4, 6)
    row.add_merge(2, 3)
    report_writer.write_row(ws, row, row_index=5)
    merged = {str(r) for r in ws.merged_cells.ranges}
    assert "D5:F5" in merged
    assert "B5:C5" in merged


def test_write_row_no_merges_by_default(ws):
    row = ReportRow(style_name="Normal")
    report_writer.write_row(ws, row, row_index=5)
    assert len(ws.merged_cells.ranges) == 0


def test_write_row_unsupported_formula_shape_raises(ws):
    row = ReportRow(style_name="Normal")
    row.set_formula(1, "=SUM(R1C1:R5C5)")  # a cross-column shape, never used, not supported
    with pytest.raises(ValueError, match="doesn't match the only supported"):
        report_writer.write_row(ws, row, row_index=1)


# ---------------------------------------------------------------------------
# write_section / write_sections
# ---------------------------------------------------------------------------

def _two_row_section(id: str) -> ReportSection:
    section = ReportSection(id)
    r1 = ReportRow(style_name="Normal")
    r1.set_value(1, f"{id}-row1")
    section.add_row(r1)
    r2 = ReportRow(style_name="Normal")
    r2.set_value(1, f"{id}-row2")
    section.add_row(r2)
    return section


def test_write_section_returns_next_row(ws):
    next_row = report_writer.write_section(ws, _two_row_section("A"), start_row=1)
    assert next_row == 3
    assert ws.cell(row=1, column=1).value == "A-row1"
    assert ws.cell(row=2, column=1).value == "A-row2"


def test_write_sections_chains_back_to_back(ws):
    sections = [_two_row_section("A"), _two_row_section("B")]
    next_row = report_writer.write_sections(ws, sections, start_row=1)
    assert next_row == 5
    assert ws.cell(row=1, column=1).value == "A-row1"
    assert ws.cell(row=2, column=1).value == "A-row2"
    assert ws.cell(row=3, column=1).value == "B-row1"
    assert ws.cell(row=4, column=1).value == "B-row2"


# ---------------------------------------------------------------------------
# apply_header_footer
# ---------------------------------------------------------------------------

def test_apply_header_footer_sets_expected_fields(ws):
    report_writer.apply_header_footer(
        ws,
        center_header="36 Elements",
        customer_name="Acme Corp",
        customer_address_line1="123 Main St",
        customer_address_line2="Springfield, IL",
        lab_header_lines=("Some Lab", "100 Industrial Pkwy", "Phone: 555-0100"),
    )
    assert ws.oddHeader.left.text == "Some Lab\n100 Industrial Pkwy\nPhone: 555-0100"
    assert ws.oddHeader.center.text == "36 Elements"
    assert "Acme Corp" in ws.oddHeader.right.text
    assert ws.oddFooter.center.text == "&F"
    assert ws.oddFooter.right.text == "Page &P"
    assert ws.page_margins.top == 0.99


def test_apply_header_footer_survives_a_real_save_and_reload(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    report_writer.apply_header_footer(
        ws,
        center_header="36 Elements",
        customer_name="Acme Corp",
        customer_address_line1="123 Main St",
        customer_address_line2="Springfield, IL",
    )
    path = tmp_path / "test.xlsx"
    wb.save(path)

    reloaded = openpyxl.load_workbook(path)
    reloaded_ws = reloaded.active
    assert reloaded_ws.oddHeader.center.text == "36 Elements"
    assert reloaded_ws.oddFooter.center.text == "&F"
    assert reloaded_ws.page_margins.top == 0.99


def test_apply_header_footer_no_lab_lines_leaves_left_header_unset(ws):
    """Deliberately left unset (not assigned an empty string) -- see the
    comment in apply_header_footer for the real openpyxl serialization bug
    this avoids."""
    report_writer.apply_header_footer(
        ws,
        center_header="TOC",
        customer_name="Acme Corp",
        customer_address_line1="",
        customer_address_line2="",
    )
    assert ws.oddHeader.left.text is None


def test_apply_header_footer_no_lab_lines_survives_save_and_reload(tmp_path):
    """Regression test for the empty-string-corrupts-neighboring-parts bug:
    confirms center/right still come back correctly when left is never set."""
    wb = openpyxl.Workbook()
    ws = wb.active
    report_writer.apply_header_footer(
        ws,
        center_header="TOC",
        customer_name="Acme Corp",
        customer_address_line1="",
        customer_address_line2="",
    )
    path = tmp_path / "test.xlsx"
    wb.save(path)

    reloaded = openpyxl.load_workbook(path)
    reloaded_ws = reloaded.active
    assert reloaded_ws.oddHeader.center.text == "TOC"
    assert reloaded_ws.oddHeader.right.text is not None
    assert "Acme Corp" in reloaded_ws.oddHeader.right.text


# ---------------------------------------------------------------------------
# apply_column_widths
# ---------------------------------------------------------------------------

def test_apply_column_widths_sets_each_listed_column(ws):
    report_writer.apply_column_widths(ws, {1: 17.14, 3: 6.57})
    assert ws.column_dimensions["A"].width == 17.14
    assert ws.column_dimensions["C"].width == 6.57


def test_apply_column_widths_leaves_unlisted_columns_alone(ws):
    report_writer.apply_column_widths(ws, {1: 17.14})
    assert "B" not in ws.column_dimensions
